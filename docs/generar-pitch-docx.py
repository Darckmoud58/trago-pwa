#!/usr/bin/env python3
"""Genera el Word comercial de TraGo (portada + tesis de venta)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "TraGo-Pitch-Comercial.docx"

INK = RGBColor(0x1A, 0x18, 0x14)
MUTED = RGBColor(0x5C, 0x5A, 0x54)
COPPER = RGBColor(0xB8, 0x6B, 0x2A)
LEAF = RGBColor(0x1C, 0x3D, 0x34)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF7, 0xF3, 0xEA)
RULE = RGBColor(0xD4, 0xC8, 0xB0)


def set_run(run, *, size=11, bold=False, color=INK, font="Calibri", italic=False):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D4C8B0")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_text(cell, text, *, size=10, bold=False, color=INK, fill=None, center=False):
    if fill:
        shade(cell, fill)
    set_cell_border(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)


def add_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("TraGo  ·  Documento comercial confidencial  ·  Septiembre 2026")
    set_run(r, size=8, color=MUTED, italic=True)

    p2 = footer.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run("Pág. ")
    set_run(r2, size=8, color=MUTED)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r3 = p2.add_run()
    r3._r.append(fld1)
    r3._r.append(instr)
    r3._r.append(fld2)
    set_run(r3, size=8, color=MUTED)


def hbar(doc, color="1C3D34"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def heading(doc, text, n=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if n == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    if n == 1:
        set_run(run, size=18, bold=True, color=LEAF, font="Calibri")
        hbar(doc, "B86B2A")
    else:
        set_run(run, size=13, bold=True, color=COPPER, font="Calibri")
    return p


def para(doc, text, *, size=11, italic=False, bold=False, color=INK, space=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, italic=italic, color=color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.75)
    run = p.add_run("•  " + text)
    set_run(run, size=11, color=INK)
    return p


def table(doc, headers, rows, col_widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, size=9, bold=True, color=WHITE, fill="1C3D34", center=True)
    for r_i, row in enumerate(rows):
        fill = "F7F3EA" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell_text(t.rows[r_i + 1].cells[c_i], val, size=9, fill=fill)
    if col_widths:
        for row in t.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TRAGO")
    set_run(r, size=48, bold=True, color=LEAF, font="Calibri")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("La promo que sí está viva — en la sucursal de a lado")
    set_run(r, size=18, color=COPPER, font="Calibri")

    hbar(doc, "B86B2A")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Documento comercial  ·  Tesis de producto, audiencia y rentabilidad"
    )
    set_run(r, size=12, italic=True, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PWA  ·  Cadenas  ·  Sucursales  ·  Geolocalización  ·  México 18+")
    set_run(r, size=11, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Septiembre 2026  ·  Uso interno y presentación a socios")
    set_run(r, size=10, color=MUTED)

    doc.add_page_break()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    add_footer(section)

    cover(doc)

    heading(doc, "Una frase")
    para(
        doc,
        "Hoy las cadenas publican promociones. Mañana el cliente llega a la sucursal y ya no aplican. TraGo vende certeza: qué promo sigue vigente, en cuál local, a cuántos metros.",
        italic=True,
        size=13,
        color=LEAF,
    )

    heading(doc, "1. El problema que resolvemos")
    para(
        doc,
        "México está lleno de “2x1”, “3x2 en botella” y “cover + mixto”. El anuncio vive en Instagram, en el volante de la cadena o en el sitio corporativo. La vigencia vive en el piso — y el piso no es uniforme.",
    )
    para(doc, "Tres fallas que se repiten cada noche:", bold=True)
    bullet(
        doc,
        "La promo es de la cadena, no del local. Polanco sí; Condesa ya se acabó. El usuario no tiene cómo saberlo sin ir.",
    )
    bullet(
        doc,
        "El contenido caduca y nadie lo baja. El post de hace 12 días sigue likeable. El cliente se mueve. El negocio pierde la venta y gana un comentario malo.",
    )
    bullet(
        doc,
        "Quien paga publicidad compra clics, no pies en la puerta. Una cadena gasta en “estamos en promo” y convierte mal porque una parte de las sucursales ya no participa, se les acabó el inventario o el gerente no activó el folio.",
    )
    para(
        doc,
        "Eso no es un problema de “más catálogo”. Es un problema de verdad local y a tiempo. El ticket nocturno (botella, mesa, six) es alto. El arrepentimiento también. Quien resuelva “¿sí o no, aquí y ahora?” se queda con la intención de compra más cara del retail de conveniencia y de la vida nocturna.",
    )

    table(
        doc,
        ["Actor", "Qué pierde hoy"],
        [
            ["Usuario", "Tiempo, pasaje, la noche, confianza"],
            ["Cadena / sucursal", "Visita frustrada, mala reseña, inventario mal empujado"],
            ["Marca (destilado, cerveza)", "Trade spend sin prueba de que la botella se movió en ese local"],
            ["App o medio genérico", "Promos “nacionales” que mienten en la esquina"],
        ],
        [5, 12],
    )

    heading(doc, "2. A quién va dirigido")
    para(
        doc,
        "TraGo no es “para todo el que sale”. Tiene dos caras que se alimentan: quien abre la PWA y quien paga por aparecer con la verdad.",
    )

    heading(doc, "A. Consumidor — demanda", 2)
    table(
        doc,
        ["Perfil", "Qué busca", "Por qué TraGo"],
        [
            [
                "Cazador de promo 18–35",
                "2x1, six, michelada, botella esta noche",
                "GPS + vigencia por sucursal, no un feed infinito",
            ],
            [
                "Grupo que sale",
                "Cover + botella, rooftop, antro",
                "Decide en 30 segundos si vale el Uber",
            ],
            [
                "Comprador de destilado",
                "3x2 mezcal / 20% whisky en licorería",
                "Ve si esa sucursal todavía lo tiene",
            ],
        ],
        [4.5, 6.5, 6],
    )
    para(doc, "Promesa al usuario: no te muevas de balde. PWA instalable; sin App Store el día uno.", italic=True)

    heading(doc, "B. Cadena y marca — quien paga", 2)
    table(
        doc,
        ["Perfil", "Dolor", "Qué le vendemos"],
        [
            [
                "Cadena de conveniencia",
                "Promo nacional que no pega igual en 200 tiendas",
                "Publicar por sucursal; ver dónde confirman o niegan",
            ],
            [
                "Licorería / cava",
                "Botella cara y fechas cortas",
                "Empujar 3x2 solo donde hay stock y tráfico",
            ],
            [
                "Antros / cantinas",
                "Cover y botella que cambian jueves vs sábado",
                "Destacar “cerca de ti” en horario nocturno",
            ],
            [
                "Marca de alcohol",
                "Paga exhibición y no sabe si el local cumplió",
                "Promo patrocinada + señal de vigencia en piso",
            ],
        ],
        [4.5, 6.5, 6],
    )
    para(
        doc,
        "No es el cliente primario el directorio de “lugares bonitos” (eso ya lo hace Google e Instagram). TraGo gana si la promo es el producto. No competimos con Rappi: competimos con el viaje en vano.",
    )

    heading(doc, "3. La solución")
    para(
        doc,
        "TraGo = promo oficial de la cadena × sucursales participantes × GPS × verdad de la calle.",
        bold=True,
        color=LEAF,
    )
    bullet(doc, "La empresa publica la promo: fechas, tipo (2x1, botella, cover, happy hour), términos, sucursales.")
    bullet(doc, "El usuario ve las que están vigentes cerca.")
    bullet(doc, "Entra al detalle y ve en cuáles locales sigue viva.")
    bullet(doc, "En el piso confirma o niega: “sigue vigente” / “ya no aplica”.")
    bullet(
        doc,
        "Si se acumulan reportes en contra, esa sucursal pasa a “en duda” o se oculta. La cadena no pierde el control de lo oficial; el piso corrige el folleto.",
    )
    para(
        doc,
        "Eso cierra la pregunta abierta con Lalo: no es o la empresa o el usuario. Son las dos capas. Sin la empresa no hay oferta legal ni marca. Sin el usuario, la vigencia es un PDF.",
    )
    para(
        doc,
        "Por qué PWA: se instala desde el navegador, geolocalización nativa, un solo código (Next.js) para web y “app”. Wrappers nativos cuando haya tracción — no antes.",
        italic=True,
    )

    heading(doc, "4. Por qué es vendible y rentable")
    para(
        doc,
        "El dinero no está en el usuario gratuito. Está en quien necesita que esa sucursal se llene esta noche.",
    )
    para(doc, "Cinco motivos para que la cadena o la marca paguen:", bold=True)
    bullet(doc, "Aparecer primero en “Cerca de ti” cuando hay intención real (GPS + horario nocturno).")
    bullet(doc, "Dejar de quemar presupuesto en sucursales donde la promo ya murió.")
    bullet(doc, "Prueba social de piso: N personas dijeron “sigue vigente” en Reforma vs 0 en Del Valle.")
    bullet(doc, "API para meter TraGo en su app, letrero digital o CRM (plan Premium).")
    bullet(doc, "Marcas de alcohol: pauta nativa sobre la promo de botella, no un banner genérico.")

    heading(doc, "Modelo de ingresos", 2)
    table(
        doc,
        ["Plan", "Quién", "Qué incluye", "Precio de lista sugerido"],
        [
            [
                "Freemium",
                "Cadena chica / cantina",
                "Perfil + 2 promos activas + anuncios de terceros",
                "$0",
            ],
            [
                "Pro",
                "Cadena con 3+ sucursales",
                "Sin anuncios propios, destacada cerca de ti, reportes de vigencia",
                "$2,900 – $6,900 MXN / mes por zona",
            ],
            [
                "Premium",
                "Grupo / marca",
                "API, sucursales ilimitadas, white-label light, export",
                "$12,000 – $25,000 MXN / mes",
            ],
            [
                "Campaña de marca",
                "Destilado / cerveza",
                "Promo patrocinada 7–15 días, slot nocturno",
                "$15,000 – $80,000 por vuelo",
            ],
            [
                "Pauta Freemium",
                "Anunciantes locales",
                "Banner en perfiles free",
                "CPM / CPC local",
            ],
        ],
        [3.2, 4.2, 5.6, 4],
    )
    para(
        doc,
        "Precios de lista para vender la idea y el piloto; no son un SAP. Se ajustan con las primeras 10 cadenas.",
        size=9,
        italic=True,
        color=MUTED,
    )

    heading(doc, "Unidad económica — piloto CDMX (conservador, 12 meses)", 2)
    table(
        doc,
        ["Palanca", "Orden de magnitud año 1"],
        [
            ["40 cadenas Pro × $4,000 / mes promedio", "~ $1.9 M MXN"],
            ["8 Premium × $15,000 / mes", "~ $1.4 M MXN"],
            ["6 vuelos de marca / trimestre", "~ $1.2 – $2.0 M MXN"],
            ["Pauta Freemium", "~ $0.3 M MXN"],
            ["Ingreso bruto indicativo", "$4.8 – $5.6 M MXN"],
            ["Costo (2–3 personas + infra + adquisición)", "$2.5 – $3.2 M MXN"],
        ],
        [10, 7],
    )
    para(
        doc,
        "El margen de contribución es positivo si el piloto cierra 25+ planes Pro en 6 meses. No hace falta ser “el Uber de las promos”. Hace falta ser el sistema de verdad de 40 cadenas que ya gastan en volante, pauta y trade. Software, no inventario. PWA: cero comisión de tienda al inicio. El contenido lo ponen las cadenas; los usuarios lo curan.",
    )

    heading(doc, "5. Competencia y cuña")
    table(
        doc,
        ["Alternativa", "Qué hace", "Dónde falla vs TraGo"],
        [
            ["Instagram / TikTok", "Alcance", "Sin vigencia por sucursal ni “aún aplica”"],
            ["Google Maps", "Dónde está el local", "No opera folios de promo; el post se pudre"],
            ["Lealtad de cadena", "Descuento propio", "Silo: no comparas cava vs antro vs tienda"],
            ["Rappi / PedidosYa", "Entrega", "Otro trabajo: “me lo traen”, no “¿voy o no?”"],
            ["Grupos de Facebook", "Ruido social", "Sin oficialidad, sin mapa, sin API"],
        ],
        [4.2, 4.5, 8.3],
    )
    para(doc, "Cuña: oficial + local + vivo. Nadie serio junta las tres.", bold=True, color=COPPER)

    heading(doc, "6. Go-to-market")
    para(doc, "Orden ejecutable, no un wish list:", bold=True)
    bullet(
        doc,
        "Oferta primero: 10–15 cadenas CDMX (licorerías, un grupo de antros, un formato conveniencia). Ellas cargan promos reales.",
    )
    bullet(
        doc,
        "Demanda en corredores: Roma–Condesa–Reforma–Polanco. Un mensaje: botella y 2x1 que todavía sirven.",
    )
    bullet(
        doc,
        "Prueba al gerente: “12 personas confirmaron vigencia anoche en esta sucursal; 9 dijeron que en la otra ya no.” Eso vende el plan Pro.",
    )
    bullet(doc, "Marcas en vuelo 2, cuando haya 20+ sucursales con tráfico.")
    para(
        doc,
        "Métrica norte: sesiones con GPS × promo abierta × sucursal vista. No “likes”.",
        italic=True,
    )

    heading(doc, "7. Producto: hoy y lo que se cobra después")
    para(doc, "Hoy (MVP)", bold=True, color=LEAF)
    bullet(doc, "Home con geolocalización (fallback Centro CDMX).")
    bullet(doc, "Promos vigentes y filtro por tipo.")
    bullet(doc, "Detalle: sucursales + voto de vigencia.")
    bullet(doc, "Capa Nocturno (botellas, cover, 2x1).")
    bullet(doc, "Perfiles de cadena Freemium / Pro / Premium y API de ejemplo.")
    bullet(doc, "PWA instalable; contenedor Docker para levantarlo igual en cualquier máquina.")
    para(doc, "90 días — para cobrar", bold=True, color=LEAF)
    bullet(doc, "Login de cadena y alta de promos / sucursales.")
    bullet(doc, "Base de datos (sale el mock).")
    bullet(doc, "Moderación de reportes y panel de vigencia por local.")
    bullet(doc, "Gate 18+ y geocerca de alcohol.")
    para(doc, "12 meses", bold=True, color=LEAF)
    bullet(doc, "Reservación / cover (take rate).")
    bullet(doc, "Push: “tu 3x2 de mezcal caduca en 4 horas a 800 m”.")
    bullet(doc, "MTY / GDL con el mismo playbook.")

    heading(doc, "8. Riesgos — cómo se contestan en una mesa")
    table(
        doc,
        ["Riesgo", "Respuesta"],
        [
            [
                "Usuarios trollan la vigencia",
                "Umbral + cuenta + la cadena reafirma lo oficial. El voto no borra el folio: lo matiza.",
            ],
            [
                "Las cadenas no cargan sucursal por sucursal",
                "Excel / API / “aplicar a todas excepto…”. En piloto lo cargamos nosotros.",
            ],
            [
                "Alcohol y 18+",
                "Gate, copy 18+, no venta de alcohol: información de promo. Aviso de privacidad desde el piloto.",
            ],
            [
                "“Google ya está”",
                "Google no cobra ni opera vigencia de 2x1 por folio de sucursal. Ese es el hueco.",
            ],
            [
                "¿Empresa o usuario?",
                "El híbrido es la tesis. Se valida en 4 semanas de piloto.",
            ],
        ],
        [5, 12],
    )

    heading(doc, "9. Guion de venta (60 segundos)")
    para(
        doc,
        "Las cadenas ya pagan por decir “tenemos promo”. El cliente llega y en la mitad de las sucursales no es cierto. TraGo es la capa que responde sí o no, aquí, ahora, con GPS, y deja que la gente en piso confirme. Cobramos a la cadena por aparecer cerca y por no lucir sucursales muertas. Al usuario le damos la noche sin viaje en vano. Es PWA, se instala ya, y el piloto en CDMX se mide en pies en la puerta — no en likes.",
        italic=True,
        size=12,
        color=LEAF,
    )

    heading(doc, "10. Pedido — siguiente paso")
    para(doc, "Para pasar de demo a piloto cobrable:")
    bullet(doc, "8–12 cadenas ancla en CDMX (aunque sea una sucursal cada una).")
    bullet(doc, "Una persona de contenido / onboarding, 4 semanas.")
    bullet(doc, "Auth + Postgres (el contenedor Docker ya levanta el producto).")
    bullet(doc, "Precio Pro de lista: $2,900 vs $4,900 según se venda por cadena o por zona.")
    para(
        doc,
        "Cierre: TraGo no pide que el mercado “descubra las salidas”. Pide que deje de mentir la distancia entre el anuncio y la caja.",
        bold=True,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run(
        "Cifras de ingreso son escenarios de lista para la conversación comercial, no un forecast auditado."
    )
    set_run(r, size=8, italic=True, color=MUTED)

    doc.save(OUT)
    print(f"OK {OUT}")


if __name__ == "__main__":
    build()
